#!/usr/bin/env python3
"""Build a 3-month / ¥120,000 quotation workbook for P1–P3."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "14213D"
GOLD = "C45C26"
INK = "1A1A1A"
MUTED = "5C5C5C"
WHITE = "FFFFFF"
CREAM = "F7F4EF"
LINE = "D9D4C8"
P1C = "1F6B4A"
P2C = "1F4E79"
P3C = "7A3E12"
OK_BG = "E8F0EC"
THIN = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
WRAP = Alignment(wrap_text=True, vertical="center")
LEFT = Alignment(wrap_text=True, vertical="center", horizontal="left")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
RIGHT = Alignment(wrap_text=True, vertical="center", horizontal="right")


def fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def font(size=11, bold=False, color=INK, name="Calibri") -> Font:
    return Font(name=name, size=size, bold=bold, color=color)


def widths(ws, mapping: dict[str, float]) -> None:
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def paint_header_row(ws, row: int, ncols: int, color: str = NAVY) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.font = font(11, True, WHITE)
        cell.fill = fill(color)
        cell.alignment = CENTER
        cell.border = THIN
    ws.row_dimensions[row].height = 28


def paint_body(ws, start: int, end: int, ncols: int, money_cols: set[int] | None = None) -> None:
    money_cols = money_cols or set()
    for r in range(start, end + 1):
        ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 0, 36)
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c)
            cell.font = font()
            cell.alignment = WRAP
            cell.border = THIN
            cell.fill = fill(CREAM if r % 2 == 0 else WHITE)
            if c in money_cols and isinstance(cell.value, (int, float)):
                cell.number_format = '"¥"#,##0'
                cell.alignment = RIGHT


def cover(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "报价摘要"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY

    ws["A1"] = "项目报价单"
    ws["A1"].font = font(22, True, NAVY)
    ws["A2"] = "海外拓客全域整合 · 官网 AI 智能助理 · 社媒自动化中枢"
    ws["A2"].font = font(13, False, MUTED)
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 22

    meta = [
        ("项目名称", "SLAM / 外贸出海智能获客与触达平台（P1–P3）"),
        ("建设周期", "3 个自然月（含联调、试点、验收）"),
        ("报价币种", "人民币 CNY"),
        ("含税总价", "¥120,000（人民币壹拾贰万元整，含税）"),
        ("技术主轴", "LangGraph 多智能体编排 + Dify 知识与工作流运营面"),
        ("交付形态", "可运行系统 + 知识库/规则配置后台 + 导出与审计 + 架构说明"),
        ("报价有效期", "自发出之日起 30 日"),
    ]
    ws["A4"] = "项"
    ws["B4"] = "内容"
    paint_header_row(ws, 4, 2)
    for i, (k, v) in enumerate(meta, start=5):
        ws.cell(i, 1, k).font = font(11, True, NAVY)
        ws.cell(i, 2, v).font = font()
        ws.cell(i, 1).border = THIN
        ws.cell(i, 2).border = THIN
        ws.cell(i, 1).fill = fill(CREAM)
        ws.cell(i, 2).fill = fill(WHITE)
        ws.cell(i, 2).alignment = LEFT
        ws.row_dimensions[i].height = 24

    ws["A13"] = "三板块报价一览"
    ws["A13"].font = font(14, True, NAVY)
    headers = ["板块", "优先级", "建设主题", "含税金额（元）", "占比", "周期重心"]
    for c, h in enumerate(headers, 1):
        ws.cell(14, c, h)
    paint_header_row(ws, 14, 6)
    rows = [
        ("P1 海外拓客全域整合", "最高", "海关发现 → 官网深挖 → LLM 清洗 → 证据化导出", 56000, "46.7%", "第 1–2 月"),
        ("P2 官网 AI 智能助理", "次高", "Dify RAG 接待 + 意图挖掘 + WhatsApp 转接 + SEO/GEO", 40000, "33.3%", "第 2–3 月"),
        ("P3 海外社媒自动化", "后续", "OAuth 多平台分发 + 语义筛客 + 风控编排中枢", 24000, "20.0%", "第 3 月"),
        ("合计", "—", "平台脊柱一次性纳入各板块，不另计", 120000, "100%", "3 个月"),
    ]
    for i, row in enumerate(rows, start=15):
        for c, val in enumerate(row, 1):
            ws.cell(i, c, val)
        ws.row_dimensions[i].height = 40
    paint_body(ws, 15, 18, 6, money_cols={4})
    for c in range(1, 7):
        ws.cell(18, c).font = font(11, True, WHITE)
        ws.cell(18, c).fill = fill(NAVY)
        ws.cell(18, c).alignment = CENTER if c != 4 else RIGHT
    ws["D18"].number_format = '"¥"#,##0'

    ws["A20"] = "架构一句话"
    ws["A20"].font = font(14, True, NAVY)
    ws["A21"] = (
        "以 LangGraph 作为「决策与工具调用脊柱」（多 Agent、检查点、条件边、证据门禁），"
        "以 Dify 作为「可运营面」（知识库、Chatflow、内容工作流、非研发人员可改），"
        "以 LiteLLM 作为统一模型网关；海关/抓取/分发等确定性能力用 GitHub 成熟开源组件装配，而不是从零堆模型。"
        "P1 产线索，P2 收询盘并做 GEO，P3 把内容与筛客回流进同一线索池。"
    )
    ws["A21"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[21].height = 78

    ws["A23"] = "付款节奏（建议）"
    ws["A23"].font = font(14, True, NAVY)
    pay_h = ["节点", "比例", "金额（元）", "触发条件"]
    for c, h in enumerate(pay_h, 1):
        ws.cell(24, c, h)
    paint_header_row(ws, 24, 4, GOLD)
    pays = [
        ("合同签订 / 开工", "40%", 48000, "启动平台脊柱、P1 海关发现与深挖骨架"),
        ("第 8 周中期验收", "40%", 48000, "P1 可导出清洗后线索；P2 挂件可基于知识库拒答"),
        ("第 12 周终验", "20%", 24000, "P2 转接与 SEO/GEO 清单完成；P3 OAuth 分发+筛客试点"),
    ]
    for i, row in enumerate(pays, start=25):
        for c, val in enumerate(row, 1):
            ws.cell(i, c, val)
        ws.row_dimensions[i].height = 32
    paint_body(ws, 25, 27, 4, money_cols={3})

    widths(ws, {"A": 28, "B": 16, "C": 52, "D": 18, "E": 12, "F": 16})


def modules(wb: Workbook) -> None:
    ws = wb.create_sheet("分板块报价")
    ws.sheet_properties.tabColor = P1C
    headers = ["板块", "工作包", "方案要点（架构级）", "开源/引擎", "人月重心", "金额（元）"]
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h)
    paint_header_row(ws, 1, 6)
    rows = [
        ("P1", "平台脊柱", "LangGraph StateGraph 多 Agent 运行时、检查点、SSE 任务总线、LiteLLM 双模型网关（推理/抽取解耦）、租户级配置与成本观测", "LangGraph · LiteLLM · FastAPI", "贯穿 3 月", 12000),
        ("P1", "海关发现层", "HS / 采购国 / 品名 / 频次驱动的进口商发现；公开源做市场验证，授权商业源做出企业名单；现有「公司→海关证据」降为补全，不再爬付费站前台", "UN Comtrade API · 本仓 customs_router · b2b-lead-hunter-skill 证据范式", "第 1 月", 16000),
        ("P1", "官网图谱深挖", "直连 → 阅读器 → 无头浏览器三级瀑布；同域多层（联系/团队/采购/Impressum）；决策人、邮箱、社媒、办公联络结构化抽取", "Crawlee-Python · Playwright · Jina Reader · contact_extractor", "第 1–2 月", 14000),
        ("P1", "LLM 清洗与意图", "货代/同行/空邮箱(MX)/活跃窗口规则引擎 + 模型分类；破产仅三档标注；A/B/C 意图；全程 evidence[] 可溯源", "LangGraph Evaluate · OpenCorporates 客户端", "第 2 月", 8000),
        ("P1", "导出与权属", "CSV/JSONL 批量导出（含证据 URL）；不开发客户 CRM/分析", "lead-hunter JSONL 产物链", "第 2 月", 6000),
        ("P2", "Dify 知识运营面", "手册/FAQ/报价边界入库；客户可自更新；Chatflow 多语接待（德/英/中）", "Dify Knowledge + Chatflow", "第 2 月", 12000),
        ("P2", "拒答门禁接待", "检索为空则不调用生成模型；禁止编型号/价格/认证；LangGraph 做工具侧硬门禁，Dify 做人可改话术", "LangGraph retrieval-or-refuse · Dify", "第 2 月", 10000),
        ("P2", "意图与转人工", "对话抽取采购需求/预算/周期；高意向 webhook；WhatsApp 走官方 Cloud API 一键深链（不收个人号密码）", "Chatwoot（可选接待台）· Meta Cloud API", "第 3 月", 8000),
        ("P2", "SEO / GEO 解耦层", "聊天脚本 defer，不改原业务逻辑；JSON-LD + sitemap + llms.txt + AI 爬虫策略；不承诺排名、不承诺必被大模型点名引用", "llms.txt 规范 · aio-surfaces · schema.org 模板", "第 2–3 月", 10000),
        ("P3", "OAuth 分发中枢", "FB/IG/X/LinkedIn/YouTube 官方 OAuth 绑定/解绑；一次提交、分平台适配；文件限额预检", "Postiz（独立部署，AGPL 隔离）", "第 3 月", 12000),
        ("P3", "语义筛客回流", "可配关键词组 + 多模态 LLM 语义（非页面 RPA）；自有频道官方字幕；筛客结果进 P1 线索池", "Dify 工作流 · LiteLLM 视觉模型 · YouTube Captions", "第 3 月", 8000),
        ("P3", "风控编排", "平台级配额、随机间隔、审计日志、一键熔断；互动以名单+深链+人工确认为闭环，不采用云手机/模拟点击", "Postiz 调度 + 自研限流审计", "第 3 月", 4000),
    ]
    for i, row in enumerate(rows, start=2):
        for c, val in enumerate(row, 1):
            ws.cell(i, c, val)
        ws.row_dimensions[i].height = 58
    paint_body(ws, 2, 13, 6, money_cols={6})
    for r, row in enumerate(rows, start=2):
        color = {"P1": P1C, "P2": P2C, "P3": P3C}[row[0]]
        ws.cell(r, 1).font = font(11, True, WHITE)
        ws.cell(r, 1).fill = fill(color)
        ws.cell(r, 1).alignment = CENTER

    ws["A15"] = "含税合计（平台脊柱已摊入 P1 首包，不重复计费）"
    ws["F15"] = 120000
    ws["F15"].number_format = '"¥"#,##0'
    for c in range(1, 7):
        ws.cell(15, c).fill = fill(NAVY)
        ws.cell(15, c).border = THIN
        ws.cell(15, c).font = font(12, True, WHITE)
        ws.cell(15, c).alignment = CENTER
    ws["A15"].alignment = LEFT
    ws["F15"].alignment = RIGHT

    ws["A17"] = "说明：金额按架构工作包切分，不按「每个页面/每条选择器」计价。网站改版若只影响通用抽取，含在三个月实施与一个月免费缺陷修复内；指定死盯站点另议。"
    ws["A17"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[17].height = 40
    widths(ws, {"A": 10, "B": 18, "C": 58, "D": 42, "E": 14, "F": 14})


def stack(wb: Workbook) -> None:
    ws = wb.create_sheet("技术栈与开源装配")
    ws.sheet_properties.tabColor = P2C
    headers = ["层级", "选型", "在方案中的角色", "复杂度说明", "来源"]
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h)
    paint_header_row(ws, 1, 5, P2C)
    rows = [
        ("编排脊柱", "LangGraph StateGraph", "多 Agent 状态机：洞察 → 关键词 → 检索 → 抽取 → 评估 →（可选）触达；检查点可恢复、条件边控制轮次", "比单次 Prompt 链高一个数量级：工具调用、记忆、停机条件都在图里", "langchain-ai/langgraph（本仓已落地）"),
        ("运营面", "Dify Chatflow / Workflow / Knowledge", "知识库更新、接待话术、内容分发适配流交给非研发；与 LangGraph 分工：Dify 管「人可改」，Graph 管「硬门禁」", "避免把运营配置写死在代码里", "langgenius/dify"),
        ("模型网关", "LiteLLM 双模型", "推理模型做 ReAct 决策，快模型做抽取/改写/多语；可热切换 OpenAI / Anthropic / GLM / MiniMax 等", "成本与供应商不绑死", "本仓 LiteLLM"),
        ("任务与API", "FastAPI + SSE", "长任务入队、进度流、导出 API；前端只提交与观测", "生产级异步，而不是浏览器里干等", "本仓 FastAPI"),
        ("海关宏观", "UN Comtrade Official API", "HS × 国家贸易量验证，回答「这市场有没有量」", "全球覆盖但是汇总，不是企业名单", "uncomtrade/comtradeapicall"),
        ("海关企业", "授权商业 API（择一）+ 本仓 router", "进口商发现主路径；router 只补公开摘要证据", "这是 P1 能「按 HS 找人」的关键采购项，费用不含在 12 万里", "ImportGenius / Volza / Panjiva 等合同 API"),
        ("抓取运行时", "Crawlee-Python + Playwright", "站图谱、代理退避、JS 动态页升无头浏览器；遵守 robots", "三级瀑布，而不是单一 scrapy 脚本", "apify/crawlee-python"),
        ("页面阅读", "Jina Reader", "中等动态页转 Markdown，降低浏览器成本", "与 Playwright 互补", "本仓 jina_reader"),
        ("线索技能范式", "b2b-lead-hunter-skill", "证据链、JSONL/CSV、质量门禁优先于条数", "产物可审计", "xiongQvQ/-b2b-lead-hunter-skill"),
        ("接待台（可选）", "Chatwoot", "人机协同收件箱、会话留存、EU 自托管更利 GDPR", "挂件后面可以接工单，而不是只有气泡", "chatwoot/chatwoot"),
        ("GEO 表面", "llms.txt + aio-surfaces + JSON-LD", "给搜索引擎和生成式引擎可读的事实面；与知识库同源", "静态站也能做，不强制迁 WordPress", "AnswerDotAI/llms-txt · Janady13/aio-surfaces"),
        ("WP 分支（若有）", "Yoast/Rank Math 二选一 + AEO God Mode", "仅当存在 WordPress 时装配；与静态站方案互斥叠加而非重复 schema", "冲突检测，避免双份结构化数据", "Yoast · AEO-God-Mode"),
        ("分发中枢", "Postiz", "官方 OAuth 多平台定时发布，不收集密码、不 RPA", "AGPL 独立进程隔离，Hunter 只调其 API", "gitroomhq/postiz-app"),
        ("语音（授权范围）", "YouTube Captions + Whisper", "自有频道走官方字幕；授权素材可转写；不做全网盗字幕产品", "STT 准确率试点标定，不预打包票", "YouTube Data API · openai/whisper"),
        ("停业信号", "OpenCorporates 客户端", "注册状态进入清洗三档，缺数据标 unknown", "不能承诺司法意义上的破产判决", "opyncorporates"),
    ]
    for i, row in enumerate(rows, start=2):
        for c, val in enumerate(row, 1):
            ws.cell(i, c, val)
        ws.row_dimensions[i].height = 52
    paint_body(ws, 2, 16, 5)

    ws["A18"] = "装配原则"
    ws["A18"].font = font(13, True, NAVY)
    ws["A19"] = (
        "1. 编排与运营分离：LangGraph 管不可破坏的门禁（拒答、合规采集、限流熔断）；Dify 管可配置的知识与流程。\n"
        "2. 开源装配优先于自研爬虫/自研发帖：Crawlee、Postiz、Chatwoot、Comtrade 官方库直接进栈。\n"
        "3. 证据优先于条数：每条线索、每次回答、每次发布都留 source / 调用审计。\n"
        "4. 商业海关 API、云主机、模型 token、住宅代理按实报销或客户自备，不包含在 12 万实施费内。"
    )
    ws["A19"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[19].height = 88
    widths(ws, {"A": 20, "B": 28, "C": 48, "D": 36, "E": 36})


def timeline(wb: Workbook) -> None:
    ws = wb.create_sheet("三月实施计划")
    ws.sheet_properties.tabColor = GOLD
    headers = ["月份", "里程碑", "P1", "P2", "P3", "验收看什么"]
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h)
    paint_header_row(ws, 1, 6, GOLD)
    rows = [
        ("第 1 月\n底座 + 发现", "M1 脊柱跑通",
         "LiteLLM 网关、LangGraph 图、海关发现最小闭环（Comtrade 验证 + 一条发现路径）、深挖瀑布接通",
         "Dify 实例与知识库结构、挂件技术选型冻结（静态站 script / WP 短代码）",
         "Postiz 独立环境评估与 OAuth 应用申请启动（不承诺当月发帖）",
         "能对样本 HS 跑出带证据的候选公司表（覆盖范围取决于是否已有商业库）"),
        ("第 2 月\n清洗 + 接待", "M2 中期验收",
         "清洗规则可配、MX、意图 A/B/C、CSV/JSONL 导出",
         "RAG 接待上线：无命中拒答、德英中、客户可更新知识库；启动 JSON-LD / llms.txt / sitemap",
         "渠道账号与权限就绪",
         "P1 导出可用；P2 用手册提问不编造型号价格"),
        ("第 3 月\n触达 + 回流", "M3 终验",
         "发现/深挖/清洗稳态，导出字段冻结",
         "高意向转 WhatsApp Cloud、基础咨询报表、SEO/GEO 清单与 Search Console 提交（不承诺排名）",
         "OAuth 一键分发试点 + 语义筛客回流 P1 池 + 风控配额/熔断/审计",
         "三板块在同一线索语义下打通；发布与筛客有日志"),
    ]
    for i, row in enumerate(rows, start=2):
        for c, val in enumerate(row, 1):
            ws.cell(i, c, val)
        ws.row_dimensions[i].height = 92
    paint_body(ws, 2, 4, 6)
    ws["A6"] = (
        "三个月是「可运营的垂直切片」，不是把全球海关企业库、全网 STT、搜索排名一次性买断。"
        "P1 企业级名单的国家广度由客户是否采购商业海关 API 决定；未采购则合同范围写清覆盖上限。"
    )
    ws["A6"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[6].height = 48
    widths(ws, {"A": 16, "B": 16, "C": 32, "D": 32, "E": 32, "F": 32})


def boundary(wb: Workbook) -> None:
    ws = wb.create_sheet("含与不含")
    ws.sheet_properties.tabColor = P3C
    headers = ["类别", "包含在 ¥120,000 内", "不包含 / 需客户自备或另计"]
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h)
    paint_header_row(ws, 1, 3, P3C)
    rows = [
        ("软件实施", "LangGraph + Dify 架构落地、三板块约定工作包、联调、文档、1 次试点口径冻结", "后续大改版、额外国家包、额外语种包"),
        ("海关数据", "接入层、查询编排、证据字段、清洗规则", "Volza/Panjiva 等商业库年费；把「全球公开企业级海关」写成可交付物"),
        ("算力与 SaaS", "架构按可替换网关设计", "LLM token、Jina/Firecrawl、代理 IP、云主机、域名证书"),
        ("官网", "挂件嵌入、SEO/GEO 清单实施（静态站手写或 WP 插件二选一路径）", "整站重做、保证 Google 首页、保证 ChatGPT 点名引用"),
        ("社媒", "官方 OAuth 分发、语义筛客、风控配额与审计", "云手机矩阵、自动赞/粉刷量、封号必解与赔偿、X 高价企业 API 超额"),
        ("合规", "只采公开页、拒答门禁、GDPR 能力（导出删除、同意位）设计", "律师出函、欧盟代表、客户侧 DPA 律师费；破解登录墙"),
        ("CRM", "标准导出，客户自行导入", "客户 CRM 定制对接、数据分析看板（需求已明确不开发）"),
        ("维护", "验收后 30 日缺陷修复", "按年运维（建议另签 15–20% / 年，不在本次 12 万内）"),
    ]
    for i, row in enumerate(rows, start=2):
        for c, val in enumerate(row, 1):
            ws.cell(i, c, val)
        ws.row_dimensions[i].height = 48
    paint_body(ws, 2, 9, 3)
    widths(ws, {"A": 16, "B": 58, "C": 52})


def terms(wb: Workbook) -> None:
    ws = wb.create_sheet("商务条款")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "商务与验收要点"
    ws["A1"].font = font(16, True, NAVY)
    bullets = [
        "总价：人民币壹拾贰万元整（¥120,000，含税）。如需专票，税率与开票信息合同约定，总价不再上浮。",
        "工期：合同生效后 3 个自然月。因客户未提供 HS 样本、知识库、社媒开发者账号或未确认托管区域造成的等待，顺延不计违约。",
        "验收：按「三月实施计划」三张里程碑表；以可运行系统与书面清单为准，不以搜索排名、封号率、海关全球覆盖率为验收项。",
        "变更：范围外需求走变更单。商业海关 API、模型与云资源由客户自备或实报实销。",
        "知识产权：为本项目交付的配置、编排图、挂件与导出结构归客户使用；LangGraph/Dify/Postiz 等开源组件遵循其原许可（Postiz 为 AGPL，独立部署隔离）。",
        "数据：客户对导出线索与对话数据拥有使用权；上游海关原始库不得转售。",
        "P3 互动：方案内为官方 API 限额下的触达编排 + 人工确认。若坚持全自动赞/粉，需书面接受平台封号风险且不纳入验收与赔偿。",
        "有效期：本报价自发出之日起 30 日。",
    ]
    for i, text in enumerate(bullets, start=3):
        ws.cell(i, 1, f"{i-2}. {text}")
        ws.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(i, 1).font = font()
        ws.row_dimensions[i].height = 48
    ws["A12"] = "签署栏"
    ws["A12"].font = font(13, True, NAVY)
    ws["A14"] = "甲方（客户）签字 / 盖章：____________________    日期：__________"
    ws["A16"] = "乙方（实施）签字 / 盖章：____________________    日期：__________"
    ws["A14"].font = font(12)
    ws["A16"].font = font(12)
    widths(ws, {"A": 120})


def copy_to(src: Path, dest: Path) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(src.read_bytes())
    print(f"wrote {dest}")


def main() -> None:
    wb = Workbook()
    wb.properties.creator = "AI Hunter"
    wb.properties.title = "海外拓客三板块报价单（12万 / 3个月）"
    cover(wb)
    modules(wb)
    stack(wb)
    timeline(wb)
    boundary(wb)
    terms(wb)

    docs = Path("/workspace/docs")
    desktop = Path.home() / "Desktop"
    artifacts = Path("/opt/cursor/artifacts")
    docs.mkdir(parents=True, exist_ok=True)
    desktop.mkdir(parents=True, exist_ok=True)

    en = "overseas-gtm-quotation-120k.xlsx"
    cn = "海外拓客三板块报价单-12万-3个月.xlsx"
    primary = docs / en
    wb.save(primary)
    print(f"wrote {primary}")
    for folder in (docs, desktop, artifacts):
        if folder == artifacts and not artifacts.exists():
            continue
        copy_to(primary, folder / en)
        copy_to(primary, folder / cn)


if __name__ == "__main__":
    main()
