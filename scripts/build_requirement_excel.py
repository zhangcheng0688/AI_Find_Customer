#!/usr/bin/env python3
"""Build the customer-facing Q&A + solution workbook."""

from __future__ import annotations

import csv
import zipfile
from html import escape
from io import StringIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "1F3A5F"
ACCENT = "C45C26"
WHITE = "FFFFFF"
LIGHT = "F7F4EF"
P1 = "1F6B4A"
P2 = "1F4E79"
P3 = "7A3E12"
OK = "1F6B4A"
COND = "8A5A00"
NO = "8B1E1E"
OK_BG = "E3F2E6"
COND_BG = "FFF4D6"
NO_BG = "FDE8E8"
HEADER_FONT = Font(name="Calibri", bold=True, color=WHITE, size=11)
TITLE_FONT = Font(name="Calibri", bold=True, color=NAVY, size=18)
SUB_FONT = Font(name="Calibri", bold=True, color=NAVY, size=13)
BODY = Font(name="Calibri", size=11, color="1A1A1A")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")
THIN = Border(
    left=Side(style="thin", color="D9D4C8"),
    right=Side(style="thin", color="D9D4C8"),
    top=Side(style="thin", color="D9D4C8"),
    bottom=Side(style="thin", color="D9D4C8"),
)


def fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def style_header(ws, ncols: int, fill_color: str = NAVY) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(1, col)
        cell.font = HEADER_FONT
        cell.fill = fill(fill_color)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False


def finish_table(ws, nrows: int, ncols: int) -> None:
    """Excel-safe autofilter covering header + all data rows."""
    if nrows >= 2 and ncols >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{nrows}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def style_rows(ws, nrows: int, ncols: int, conclusion_col: int | None = None) -> None:
    for r in range(2, nrows + 1):
        ws.row_dimensions[r].height = 72
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c)
            cell.font = BODY
            cell.alignment = WRAP
            cell.border = THIN
            cell.fill = fill(LIGHT if r % 2 == 0 else WHITE)
        if conclusion_col:
            val = str(ws.cell(r, conclusion_col).value or "")
            if val.startswith("可承诺") or val.startswith("一期") or val.startswith("二期") or val.startswith("三期") or val == "可直接采用":
                ws.cell(r, conclusion_col).fill = fill(OK_BG)
                ws.cell(r, conclusion_col).font = Font(name="Calibri", size=11, color=OK, bold=True)
            elif "有条件" in val or "试点" in val or "采购" in val:
                ws.cell(r, conclusion_col).fill = fill(COND_BG)
                ws.cell(r, conclusion_col).font = Font(name="Calibri", size=11, color=COND, bold=True)
            elif "不能" in val or "不做" in val:
                ws.cell(r, conclusion_col).fill = fill(NO_BG)
                ws.cell(r, conclusion_col).font = Font(name="Calibri", size=11, color=NO, bold=True)


def set_widths(ws, widths: dict[str, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def add_hyperlink(cell, url: str) -> None:
    # Keep URL as plain text. Cell hyperlinks make some WPS/Excel builds refuse the file.
    cell.value = url
    cell.font = Font(name="Calibri", size=11, color="0B57D0", underline="single")


def cover_sheet(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "封面说明"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "海外拓客 / 官网 AI 助理 / 社媒自动化 — 需求问答与解决方案"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "给客户与内部对齐用 · 含必问问题答复、功能点方案、GitHub 开源选型、SEO/GEO 落地、分期计划"
    ws["A2"].font = Font(name="Calibri", size=12, color="4A4A4A")
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 22

    blocks = [
        (4, "这份表里有什么",
         "① 必问问题答复（需求表原 24 题，可直接转发客户）\n"
         "② 12 条核心功能点：做不做、怎么做、用哪套开源\n"
         "③ GitHub 开源选型（海关、爬虫、接待、SEO/GEO、社媒、STT）\n"
         "④ SEO 与 GEO 专项落地（静态 slam.systems 也能做，不绑死 WordPress）\n"
         "⑤ 一期/二期/三期怎么推进，以及客户需要先拍板的事项"),
        (6, "结论先看",
         "P1 海关拓客：在现有 AI Hunter 上演进，但「按 HS 找全球进口商」必须买授权海关库；公开源以美线提单为主。\n"
         "P2 官网助理 + SEO/GEO：能做好。接待用 RAG 挂件（WP 或静态站都能挂）；SEO/GEO 用 JSON-LD + sitemap + llms.txt + AI 爬虫放行，不承诺搜索排名、不承诺一定被 ChatGPT 引用。\n"
         "P3 社媒：发帖用官方 OAuth（可自托管 Postiz）；自动赞/粉不做。"),
        (8, "依据",
         "本仓库 AI Hunter 代码与测试（P1 相关 215 通过；全量 701 通过 / 1 失败在邮件 e2e）。\n"
         "上游 GitHub：xiongQvQ/AI_Find_Customer、b2b-lead-hunter-skill。\n"
         "开源选型检索日期：2026-08-19。"),
        (10, "图例",
         "绿色「可承诺 / 一期可做」= 可写进方案；黄色「有条件承诺」= 要采购海关库、试点或客户拍板；红色「不能承诺 / 不做」= 自动赞粉、封号必解、全球公开企业级海关、搜索排名保证"),
    ]
    for title_row, title, body in blocks:
        ws.cell(title_row, 1, title).font = SUB_FONT
        cell = ws.cell(title_row + 1, 1, body)
        cell.alignment = WRAP
        cell.font = BODY
        ws.row_dimensions[title_row + 1].height = 72

    ws["A13"] = "工作表"
    ws["B13"] = "内容"
    ws["A13"].font = HEADER_FONT
    ws["A13"].fill = fill(NAVY)
    ws["B13"].font = HEADER_FONT
    ws["B13"].fill = fill(NAVY)
    nav = [
        ("必问问题答复", "需求表 24 道必问问题 + 可给客户的答复 + 后续方案"),
        ("功能点与解决方案", "12 条核心功能：现状、开源组合、交付结论"),
        ("GitHub开源选型", "检索到的开源项目、许可证、用来解决哪条需求、采用方式"),
        ("SEO与GEO落地", "针对 slam.systems 静态站的 SEO/GEO 清单（不改原业务逻辑）"),
        ("分期实施计划", "一期 P1 / 二期 P2+SEO / 三期 P3 发帖筛客"),
        ("客户拍板清单", "开工前必须书面确认的事项"),
    ]
    for i, (name, desc) in enumerate(nav, start=14):
        ws.cell(i, 1, name).font = Font(name="Calibri", bold=True, color=NAVY)
        cell = ws.cell(i, 2, desc)
        cell.alignment = WRAP
        cell.font = BODY
        ws.row_dimensions[i].height = 22

    set_widths(ws, {"A": 28, "B": 72, "C": 20})


def qa_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("必问问题答复")
    headers = ["编号", "优先级", "模块", "原问题", "给客户的答复", "后续解决方案（怎么落地）", "采用的GitHub开源", "结论"]
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)
    style_header(ws, len(headers), NAVY)

    rows = [
        ["Q1", "P1", "海关数据", "覆盖哪些数据源/国家？更新周期？",
         "公开源：美线提单有企业级进口商（约T+1～T+7）；UN Comtrade / EU COMEXT 是全球HS×国家汇总，没有企业名单。全球企业级名单必须购买 ImportGenius / Panjiva / Volza 等授权API，周期跟供应商合同。现网只能给已知公司补公开摘要证据。",
         "一期接通：① UN Comtrade 官方Python库做市场验证；② 美线公开提单做企业名单试点；③ 客户选定目标国后采购1家商业海关API做发现层。现有 customs_router 只保留为「已知公司补证据」，不再爬付费站前台。",
         "uncomtrade/comtradeapicall（MIT）\nxiongQvQ/AI_Find_Customer customs_router\nxiongQvQ/-b2b-lead-hunter-skill（海关当校验）",
         "有条件承诺（要买库）"],
        ["Q2", "P1", "爬虫", "能否处理JS动态站？反爬/IP策略？",
         "一部分能。现网是Jina Reader；Playwright有对比脚本但未进生产。方案：直连→Jina→无头浏览器。遵守robots.txt、低并发、可选住宅代理。验证码/登录墙跳过，不破解。",
         "把 apify/crawlee-python 接到 LeadExtract：HTTP失败再升Playwright；同域深度2–3层；采购/团队/Impressum优先。代理与退避用Crawlee自带。不引入验证码打码服务。",
         "apify/crawlee-python（Apache-2.0）\nfirecrawl/firecrawl（AGPL，可作Jina替代，注意许可证）\n本仓库 jina_reader + Playwright demo",
         "可承诺"],
        ["Q3", "P1", "维护", "目标站改版后维护周期/成本？",
         "抽取以通用结构+LLM为主，不为每家写死XPath。常规改版含年维护。指定死盯站点另立项。海关商业源改版由供应商承担。",
         "规则（邮箱/电话/社媒正则）+ LLM抽取双轨。Crawlee不绑CSS选择器。年维护含通用适配；单站XPath外包另行报价。",
         "本仓库 contact_extractor（已有测试）\ncrawlee-python",
         "可承诺"],
        ["Q4", "P1", "数据权属", "数据能否导出？所有权归谁？",
         "能导出CSV/JSONL（含证据URL）。线索归客户自用。商业海关原始库按上游协议不转售。CRM/分析客户自己做。",
         "补 GET /api/v1/hunts/{id}/export。参考 b2b-lead-hunter-skill 的 JSONL 产物链。网页缓存定期删除。",
         "xiongQvQ/-b2b-lead-hunter-skill（MIT，JSONL/CSV导出范式）\n本仓库 hunt_store（现为JSON落盘，缺CSV接口）",
         "可承诺"],
        ["Q5", "P1", "清洗", "破产/停业怎么判定？规则能否自定义？",
         "不能100%判破产。三档：注册状态 / 很可能停业 / 未知。货代、同行、空邮箱（只做语法+MX，不做SMTP探测）、进口频次均可配。意图A/B/C。",
         "规则引擎配置化 + LLM分类。停业信号接 OpenCorporates 公开API（有配额）。邮箱沿用现有MX校验，禁止SMTP RCPT。",
         "openc/openc-schema、pjryan126/opyncorporates\n本仓库 email_verifier（测试已过）",
         "可承诺"],
        ["Q6", "P1", "合规", "是否只采公开信息？封IP责任？",
         "只采公开页。封禁导致缺条不算系统故障。不提供绕过验证码。非公开采集默认拒绝。",
         "爬取前读robots.txt；合同写明公开信息范围与免责。Crawlee遵守crawl delay。",
         "本仓库 README 合规口径 + crawlee robots支持",
         "可承诺"],
        ["Q7", "P2", "嵌入", "WordPress好不好嵌？部署多久？",
         "一段JS/短代码即可。当前slam.systems是静态站不是WP，同一脚本两种宿主都能挂。嵌入小时级；耗时在知识库和多语校对。",
         "优先：自研/二次开发Shadow DOM挂件，静态站直接插script。若客户另有WP，可评估 wpgaurav/alpha-chat（RAG+无检索则拒答）或Chatwoot网站小部件。不逼客户迁WP。",
         "wpgaurav/alpha-chat（WP RAG，无命中不调用LLM）\nchatwoot/chatwoot（3.5万星，全渠道+网站插件）\nlangchain-ai/langgraph（本仓库已用）",
         "可承诺"],
        ["Q8", "P2", "知识库", "客户能否自己更新知识库/话术？",
         "能。后台上传PDF/Markdown、勾选官网再抓、编辑拒答规则，不改代码。",
         "知识库后台：上传+URL再抓取+重建索引。若用Chatwoot/Dify类，用其文档库；若跟Hunter同栈，用现有PDF/Docx解析器。",
         "本仓库 pdf_parser / docx_parser\nchatwoot 知识库 或 Dify（Qube方案参考，非本仓）",
         "可承诺"],
        ["Q9", "P2", "语言/幻觉", "支持哪些语言？能否锁死知识库？",
         "首期德/英/中。检索不到就转人工，禁止编型号/价格/认证。硬规则，不只靠prompt。",
         "采用 alpha-chat 同类策略：检索chunks为空则直接fallback，不调用生成模型。LangGraph节点加拒答审计。",
         "wpgaurav/alpha-chat 的 empty-chunk fallback\n本仓库 LiteLLM + LangGraph",
         "可承诺"],
        ["Q10", "P2", "GDPR", "GDPR？数据归谁？",
         "对话和线索归客户。建议EU部署+DPA+同意+30天可删可导。不拿对话训基础模型。",
         "EU区主机；Chatwoot可自托管满足数据驻留。Cookie横幅只在启用统计时出现。",
         "chatwoot/chatwoot（可自托管）",
         "可承诺（托管区需确认）"],
        ["Q11", "P2", "SEO/GEO", "SEO/GEO具体做什么？「不改原功能」能否写进合同？",
         "聊天脚本defer，不改下单/表单。SEO：title/description、FAQ/Product JSON-LD、sitemap、hreflang。GEO：llms.txt + AI爬虫放行 + 与知识库同源的可引用事实。不保证搜索排名，不保证被ChatGPT引用。「不改原业务功能」可按清单入合同。",
         "slam.systems是静态HTML，不装Yoast。直接改现有页面head：Organization/Product/FAQ JSON-LD、sitemap.xml、robots.txt放行GPTBot等、根目录llms.txt。详见「SEO与GEO落地」表。WP场景才用Yoast/RankMath + AEO God Mode。",
         "AnswerDotAI/llms-txt（规范）\nfirecrawl/create-llmstxt-py、Janady13/aio-surfaces\nYoast/wordpress-seo（仅WP）\nAEO-God-Mode/aeo-god-mode（仅WP）\nJayHoltslander/Structured-Data-JSON-LD",
         "可承诺（不承诺排名/必被引用）"],
        ["Q12", "P2", "改版验收", "官网改版维护费？如何验收优化？",
         "组件与皮肤分离，大改版多半只调CSS。验收用Search Console索引、富结果、Core Web Vitals，不承诺排名。GEO验收看AI爬虫日志与llms.txt是否被抓取，仍不保证引用。",
         "SEO清单可回滚（git）。AEO God Mode有爬虫日志（仅WP）。静态站用服务器日志筛GPTBot/ClaudeBot。",
         "AEO-God-Mode/aeo-god-mode（WP爬虫日志）\n本仓库 slam-systems-site git 可回滚",
         "可承诺（不承诺排名）"],
        ["Q13", "P3", "分发", "是否OAuth？不共享密码？",
         "是。Facebook/Instagram（Meta）、LinkedIn、YouTube、X全部官方OAuth，不收集平台密码。",
         "自托管 Postiz（3.4万星，AGPL）做OAuth绑定与定时发布，覆盖需求中的FB/IG/X/LinkedIn/YouTube。注意AGPL：对客户交付需开源或独立进程隔离+网络API。",
         "gitroomhq/postiz-app（AGPL-3.0，官方OAuth，不爬不模拟登录）",
         "可承诺"],
        ["Q14", "P3", "分发", "能否一键解绑？",
         "能。撤销本系统token，并提示在平台安全页同时解除授权。",
         "Postiz/自研绑定页提供Disconnect，删除本地refresh token。",
         "gitroomhq/postiz-app",
         "可承诺"],
        ["Q15", "P3", "分发", "YouTube是否官方API？",
         "是。上传走YouTube Data API videos.insert（可恢复上传），不模拟登录。",
         "经Postiz的YouTube集成或自写videos.insert。文件走可恢复上传。",
         "gitroomhq/postiz-app YouTube集成",
         "可承诺"],
        ["Q16", "P3", "分发", "有无文件限制？",
         "有，完全跟随各官方API。YouTube约12小时或256GB（还受账号权限限制）。超限发布前拦截。",
         "发布前按平台校验时长/体积/编码，失败可见原因。",
         "各平台官方文档 + Postiz settings",
         "可承诺"],
        ["Q17", "P3", "筛客", "用哪个LLM？",
         "与Hunter同一套LiteLLM网关。默认可配gpt-4o-mini+gpt-4o，可换Anthropic/GLM/MiniMax。图文筛客用带视觉的模型。不绑死一家。",
         "筛客服务复用 backend/tools/llm_client.py。",
         "本仓库 LiteLLM（已支持OpenAI/Anthropic/OpenRouter/Groq/GLM/Moonshot/MiniMax）",
         "可承诺"],
        ["Q18", "P3", "筛客", "YouTube STT语言与准确率？",
         "自己的频道：官方Captions API。他人视频：官方API不能任意拉字幕。筛客优先标题/描述/评论。若要对公开音频做识别，需另采合规STT，准确率必须试点后给数字。",
         "自有频道走YouTube Captions。Whisper仅用于客户授权的音视频或自有素材，不用于大规模破解他人视频。不把yt-dlp偷字幕写成产品能力。",
         "openai/whisper（MIT，自有音频）\nYouTube Data API captions（自有频道）",
         "有条件承诺（需试点）"],
        ["Q19", "P3", "筛客", "关键词组能否自定义？",
         "能。多组、包含/排除、按平台开关。",
         "配置表+LLM语义二次过滤，而不是页面RPA关键词。",
         "本仓库 KeywordGen 可复用词表结构",
         "可承诺"],
        ["Q20", "P3", "筛客", "平台改UI后维护费？",
         "走官方API，不依赖页面CSS。API变更含年维护。不做RPA，故无「改UI就要重写脚本」的费用项。",
         "Postiz与自研都只调官方端点。",
         "gitroomhq/postiz-app（文档明确不爬不模拟）",
         "可承诺"],
        ["Q21", "P3", "风控", "底层是API还是RPA/云手机？",
         "官方API。不做云手机、不做模拟点击。",
         "合同写死技术路径=OAuth API。Postiz托管版同样声明官方OAuth。",
         "gitroomhq/postiz-app",
         "可承诺"],
        ["Q22", "P3", "风控", "账号有什么风险？",
         "即使用官方API，超量/重复/垃圾内容仍可能被限流或封号。自动赞/粉风险最高，故不提供。官方发帖风险低于RPA，但不能为零。",
         "发帖频率上限+内容审核队列。不接任何自动关注/点赞库。",
         "无（明确不采用RPA/云手机类仓库）",
         "可承诺（说明风险）"],
        ["Q23", "P3", "风控", "各平台每日互动上限能否自定义？",
         "发帖频率可按平台配置。自动赞/粉不做，因此无自动互动日限额。若做「名单+深链提醒销售去点」，提醒条数可配。",
         "Postiz调度间隔+自研限流。互动改为人工。",
         "gitroomhq/postiz-app 调度",
         "不能承诺自动互动"],
        ["Q24", "P3", "风控", "封号后如何补救？",
         "提供审计日志、立即停发、解绑、申诉材料（时间、API调用、内容ID）。不承诺能解封，不承担封号赔偿。建议商务号+备份账号。",
         "所有发布写审计表。一键熔断开关。",
         "自研审计 + Postiz操作日志",
         "不能承诺解封"],
    ]
    for r, row in enumerate(rows, 2):
        for c, val in enumerate(row, 1):
            ws.cell(r, c, val)
    style_rows(ws, 1 + len(rows), len(headers), conclusion_col=8)
    set_widths(ws, {"A": 8, "B": 8, "C": 12, "D": 28, "E": 42, "F": 42, "G": 36, "H": 22})
    finish_table(ws, 1 + len(rows), len(headers))


def features_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("功能点与解决方案")
    headers = ["编号", "优先级", "核心功能点", "现状（本仓库/测试）", "后续解决方案", "GitHub开源组合", "SEO/GEO是否覆盖", "结论"]
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)
    style_header(ws, len(headers), P1)
    rows = [
        ["F1", "P1", "按HS/采购国/品名/频次找活跃进口商",
         "customs_router测试通过，但是「公司→海关证据」，没有HS→进口商列表。",
         "新增Importer Discovery：Comtrade验证市场 → 美线公开提单试点 → 授权商业API出企业名单 → 交给现有LeadExtract深挖。",
         "uncomtrade/comtradeapicall\n本仓 customs_router / lead_extract_agent",
         "不适用", "一期可做（需采购商业库才能全球企业级）"],
        ["F2", "P1", "官网多层/动态/采购页/团队页，挖决策人邮箱社媒办公联系方式",
         "Jina+联系页最多3条；电话/社媒/WhatsApp测试已过。缺生产级Playwright、采购页关键词。",
         "Crawlee站图谱：contact/about/team/imprint/procurement/rfq，深度2–3，单站上限12。动态页升Playwright。",
         "apify/crawlee-python\n本仓 contact_extractor / jina_reader",
         "不适用", "一期可做"],
        ["F3", "P1", "LLM清洗：货代/破产停业/空邮箱/同行 + 意图评级",
         "已有fit/contactability/priority，会排除同类制造商。MX校验测试已过。无货代/破产专用测试。",
         "可配置规则+LLM分类；OpenCorporates补注册状态；空邮箱只做MX。",
         "本仓 email_verifier + assess_lead_fit\nopyncorporates",
         "不适用", "一期可做"],
        ["F4", "P1", "客户自己做CRM和分析，我们不开发",
         "Hunt JSON落盘，没有CSV导出接口。",
         "补export接口，不接客户CRM。",
         "b2b-lead-hunter-skill 导出范式",
         "不适用", "可承诺"],
        ["F5", "P2", "7×24多语接待，基于手册，禁止幻觉",
         "本仓库无客服代码。",
         "Shadow DOM RAG挂件；无检索命中不调用LLM。WP可用alpha-chat；静态站同一套API。复杂工单可后接Chatwoot。",
         "wpgaurav/alpha-chat\nchatwoot/chatwoot\nlangchain-ai/langgraph",
         "挂件defer，不挡爬虫", "二期可做"],
        ["F6", "P2", "从访客行为挖采购需求/预算/周期",
         "无。",
         "以对话抽取为主。全站热力图需Cookie同意，不默认隐匿追踪。",
         "接待组件会话日志 + LLM抽取（本仓llm_client）",
         "统计脚本需同意后加载", "对话抽取可承诺"],
        ["F7", "P2", "高意向同步销售，一键WhatsApp",
         "能抽取wa.me链接（测试过），无发送。",
         "高意向webhook + WhatsApp官方Cloud API。不收个人号密码。Evolution API仅作自托管备选，优先Meta Cloud。",
         "Meta Cloud API\nevolution-foundation/evolution-api（备选，注意官方政策）\nchatwoot WhatsApp通道",
         "不适用", "二期可做"],
        ["F8", "P2", "咨询与转化基础报表",
         "无。",
         "会话数/拒答数/高意向/转人工。不做CRM分析。",
         "chatwoot报表 或 自研计数",
         "不适用", "二期可做"],
        ["F9", "P2", "轻量SEO/GEO，不影响原业务逻辑",
         "slam.systems为静态HTML，已有title/description，缺JSON-LD、sitemap、llms.txt、AI爬虫策略。",
         "见「SEO与GEO落地」表：结构化数据+sitemap+llms.txt+robots放行AI爬虫。聊天组件与SEO解耦。不装必须WP的Yoast。可合同约定不改原功能。",
         "AnswerDotAI/llms-txt\naio-surfaces / create-llmstxt-py\nStructured-Data-JSON-LD\n（WP才用Yoast / AEO God Mode）",
         "本条就是SEO/GEO，可落地", "可承诺（不承诺排名/必被引用）"],
        ["F10", "P3", "一键同步发图文视频到FB/IG/X/LinkedIn/YouTube",
         "只会抽取社媒URL，不能发帖。",
         "自托管Postiz：OAuth绑定、定时、多平台适配。「同步」=一次提交分平台格式化，不是像素级同一条。AGPL需隔离部署。",
         "gitroomhq/postiz-app（34859星，AGPL-3.0）",
         "发帖内容可带站点规范URL，利于GEO引用", "三期可做"],
        ["F11", "P3", "全域筛客：图文+YouTube语音，LLM语义而非RPA",
         "无监听代码。",
         "公开搜索API+标题描述评论+LLM语义。自有频道官方字幕；Whisper仅用于授权音频。",
         "本仓 LiteLLM\nopenai/whisper（自有音频）\nYouTube Captions API",
         "不适用", "三期可做（STT有条件）"],
        ["F12", "P3", "对高意向用户自动赞/粉，随机间隔风控",
         "无。",
         "明确不做。改为名单+深链，销售人工点。GitHub上各类「auto follow」工具一律不采用。",
         "不采用任何自动关注/点赞仓库",
         "不适用", "不能承诺 / 不做"],
    ]
    for r, row in enumerate(rows, 2):
        for c, val in enumerate(row, 1):
            ws.cell(r, c, val)
    style_rows(ws, 1 + len(rows), len(headers), conclusion_col=8)
    set_widths(ws, {"A": 8, "B": 10, "C": 28, "D": 36, "E": 42, "F": 34, "G": 22, "H": 24})
    finish_table(ws, 1 + len(rows), len(headers))


def github_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("GitHub开源选型")
    headers = ["类别", "项目", "GitHub", "Stars(约)", "许可证", "解决哪条需求", "能否解决SEO/GEO", "采用方式", "风险/限制", "建议"]
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)
    style_header(ws, len(headers), ACCENT)
    rows = [
        ["P1 线索主链", "AI Hunter / AI_Find_Customer", "https://github.com/xiongQvQ/AI_Find_Customer", "151", "MIT",
         "F1–F4 搜索、抽线索、海关证据、邮件", "不负责SEO", "本仓库已是其定制分支，作为P1脊柱", "海关是校验不是发现", "采用（已在用）"],
        ["P1 线索技能", "b2b-lead-hunter-skill", "https://github.com/xiongQvQ/-b2b-lead-hunter-skill", "17", "MIT",
         "证据链、JSONL/CSV导出、合规发送门禁", "不负责", "导出与质量门禁范式复用", "同样不是HS发现引擎", "采用（范式）"],
        ["P1 海关宏观", "UN Comtrade API Python", "https://github.com/uncomtrade/comtradeapicall", "154", "MIT",
         "Q1 全球HS×国家汇总、验证市场是否有量", "不负责", "Importer Discovery第一层", "无企业名单", "采用"],
        ["P1 停业校验", "OpenCorporates 相关", "https://github.com/pjryan126/opyncorporates", "—", "开源客户端",
         "Q5 注册状态", "不负责", "有配额的公开API，失败则标unknown", "不是海关、覆盖不均", "有条件采用"],
        ["P1 爬虫", "Crawlee Python", "https://github.com/apify/crawlee-python", "9444", "Apache-2.0",
         "Q2/F2 JS动态、代理、站图谱", "不负责", "替换/增强现有Jina瀑布的浏览器层", "仍须遵守robots，不打码", "采用"],
        ["P1 抓取备选", "Firecrawl", "https://github.com/firecrawl/firecrawl", "169416", "AGPL-3.0",
         "动态页Markdown、地图网站、生成llms.txt", "create-llmstxt-py可辅助GEO", "可替换Jina；AGPL要用网络API隔离或接受开源义务", "AGPL传染；SaaS有费用", "备选"],
        ["P2 接待WP", "Alpha Chat", "https://github.com/wpgaurav/alpha-chat", "0+", "开源/GPL向",
         "Q7–Q9 RAG、无命中不调LLM、Shadow DOM", "不改主题SEO", "若客户有WP可试点；静态站不直接装", "星少、需代码审计", "WP场景试点"],
        ["P2 接待全渠道", "Chatwoot", "https://github.com/chatwoot/chatwoot", "35978", "开源可自托管",
         "Q7/Q10/F5/F7/F8 网站插件、WhatsApp、报表、EU自托管", "网站插件可defer", "二期工单/转人工/WhatsApp通道", "偏客服台，RAG需另接知识库", "采用（接待台）"],
        ["P2 编排", "LangGraph", "https://github.com/langchain-ai/langgraph", "40008", "MIT",
         "本仓库已用；接待拒答流可同栈", "不负责", "继续用", "无", "已在用"],
        ["SEO WP", "Yoast SEO", "https://github.com/Yoast/wordpress-seo", "1983", "GPL",
         "Q11 WP站点title/schema/sitemap", "能做传统SEO，GEO有限", "仅当官网是WP时安装；与AEO插件并存", "slam.systems当前不是WP", "WP才用"],
        ["SEO WP", "Rank Math", "https://github.com/RankMath/seo-by-rank-math", "134", "GPL",
         "Q11 WP schema/sitemap，与Yoast二选一", "传统SEO", "不要和Yoast同时出两套schema", "同样依赖WP", "WP才用，与Yoast二选一"],
        ["GEO/AEO WP", "AEO God Mode", "https://github.com/AEO-God-Mode/aeo-god-mode", "19", "GPL-2.0",
         "Q11–Q12 FAQ/Product schema、llms.txt、18种AI爬虫放行、冲突检测、爬虫日志", "专门做AEO/GEO，声明与Yoast/RankMath共存", "仅WP。静态站把同等能力手写进HTML", "不保证被ChatGPT引用", "WP采用；静态站等价手写"],
        ["GEO 规范", "llms.txt spec", "https://github.com/AnswerDotAI/llms-txt", "2573", "Apache-2.0",
         "Q11 根目录llms.txt让模型读站点地图", "GEO核心文件", "在slam.systems根目录放置llms.txt / llms-full.txt", "是规范不是排名魔法", "采用"],
        ["GEO 生成", "Firecrawl create-llmstxt-py", "https://github.com/firecrawl/create-llmstxt-py", "—", "随Firecrawl",
         "从现有页面生成llms.txt摘要", "辅助GEO", "一次性生成后人工审，再提交进slam-systems-site", "依赖Firecrawl/OpenAI密钥", "采用（生成稿）"],
        ["GEO 静态", "aio-surfaces", "https://github.com/Janady13/aio-surfaces", "—", "开源工具包",
         "从配置生成llms.txt、aeo.json、entity.json（Schema @graph）", "专为静态+动态站的AI引用面", "很适合slam.systems这种静态站，不改业务JS", "需人工维护事实清单", "采用（静态GEO首选）"],
        ["SEO 数据", "Structured-Data-JSON-LD", "https://github.com/JayHoltslander/Structured-Data-JSON-LD", "—", "示例集",
         "Organization/Product/FAQ JSON-LD模板", "SEO富结果+GEO可解析事实", "抄模板进现有HTML head，不改交互", "要按真实产品改，禁止编造评分", "采用（模板）"],
        ["P3 发帖", "Postiz", "https://github.com/gitroomhq/postiz-app", "34859", "AGPL-3.0",
         "Q13–Q16/F10 OAuth、FB/IG/X/LinkedIn/YouTube定时发布", "内容带规范URL有利引用", "自托管；AGPL用独立服务+API对接Hunter", "AGPL传染；官方OAuth仍有平台配额", "采用（独立部署）"],
        ["P3 发帖CLI", "Postiz Agent", "https://github.com/gitroomhq/postiz-agent", "—", "随Postiz",
         "程序化发帖", "不负责", "三期自动化接口", "依赖Postiz", "采用"],
        ["P3 STT", "OpenAI Whisper", "https://github.com/openai/whisper", "107601", "MIT",
         "Q18 自有/授权音频转写", "不负责", "仅授权素材；不作为盗字幕产品", "他人视频合规边界", "有条件采用"],
        ["P3 WhatsApp备选", "Evolution API", "https://github.com/evolution-foundation/evolution-api", "—", "开源",
         "F7 自托管WhatsApp", "不负责", "优先Meta Cloud；这个仅备选", "非官方协议有封号风险", "谨慎备选，默认不用"],
        ["P3 明确不用", "各类 auto like/follow / 云手机", "—", "—", "—",
         "F12 自动赞粉", "不负责", "全部拒绝", "违反平台ToS", "不采用"],
    ]
    for r, row in enumerate(rows, 2):
        for c, val in enumerate(row, 1):
            ws.cell(r, c, val)
        if str(row[2]).startswith("http"):
            add_hyperlink(ws.cell(r, 3), row[2])
    style_rows(ws, 1 + len(rows), len(headers), conclusion_col=10)
    set_widths(ws, {"A": 16, "B": 22, "C": 42, "D": 12, "E": 14, "F": 32, "G": 22, "H": 32, "I": 24, "J": 16})
    finish_table(ws, 1 + len(rows), len(headers))


def seo_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("SEO与GEO落地")
    headers = ["序号", "层", "动作", "是否改原业务逻辑", "开源/规范", "在slam.systems怎么做", "验收方式", "承诺边界"]
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)
    style_header(ws, len(headers), P2)
    rows = [
        ["S1", "传统SEO", "补全每页独特title、meta description、OG图", "否，只改head",
         "现有静态HTML", "在index/company/product各页写准确德/英描述，不堆砌关键词", "Search Console 索引覆盖", "可做；不承诺排名"],
        ["S2", "传统SEO", "生成sitemap.xml并提交", "否",
         "静态文件即可", "列出全部正式URL；发布流程写入slam-systems-site", "GSC sitemap状态=成功", "可做"],
        ["S3", "传统SEO", "Organization + Product + FAQ JSON-LD", "否，只加script type=application/ld+json",
         "JayHoltslander/Structured-Data-JSON-LD 模板；schema.org", "公司页Organization（瑞士主体+深圳研发事实以官网为准）；产品页Product（型号来自现有文案，禁止编价格评分）", "Google富结果测试通过", "可做；禁止虚构星级"],
        ["S4", "传统SEO", "多语hreflang（若上英/中镜像）", "否",
         "Google Search Central", "先维持德文主站；若加语言再加link rel=alternate", "GSC国际定位", "有语言版才做"],
        ["S5", "性能SEO", "聊天挂件defer/async，不挡首屏", "否，不改现有导航/表单",
         "alpha-chat Shadow DOM思路", "script放body末尾defer；失败不影响页面", "CWV不回退", "可合同写「不改原功能」"],
        ["S6", "GEO/AEO", "根目录llms.txt + llms-full.txt", "否",
         "AnswerDotAI/llms-txt 规范；firecrawl/create-llmstxt-py 或 aio-surfaces 生成", "写清公司是谁、产品线、手册入口、联系方式；与官网事实一致", "文件可公开访问；服务器日志可见GPTBot", "可做；不保证模型引用"],
        ["S7", "GEO/AEO", "aeo.json 原子问答（<500字事实）", "否",
         "Janady13/aio-surfaces", "从产品手册抽20–40条可引用事实（认证、应用场景、不编交期）", "人工审阅后上线", "可做"],
        ["S8", "GEO/AEO", "entity.json Schema @graph（组织标识对齐）", "否",
         "aio-surfaces", "同一套公司名、官网、地址，避免知识库与页面互相矛盾", "schema.org校验", "可做"],
        ["S9", "GEO/AEO", "robots.txt 放行AI爬虫（GPTBot、ClaudeBot、PerplexityBot、Google-Extended等）", "否",
         "AEO God Mode 的18爬虫名单（WP插件本身不装，只复用名单）", "在静态robots.txt Allow；若客户要拒训数据则反向Disallow并书面确认", "日志出现对应UA", "可做，策略需客户选「允许被引用」还是「拒绝训练」"],
        ["S10", "GEO/AEO", "知识库与页面同源", "否",
         "接待RAG只索引官网+已审手册", "禁止聊天回答与JSON-LD/llms.txt冲突，这是GEO的关键", "抽检问答与页面一致", "可做"],
        ["S11", "WP分支", "若未来迁WordPress", "主题功能保持，只加插件",
         "Yoast或Rank Math 二选一 + AEO God Mode（冲突检测避免双份schema）", "现在不迁。slam.systems继续静态。", "插件健康仪表盘", "暂不实施"],
        ["S12", "明确不做", "保证Google首页 / 保证ChatGPT点名引用 / 改导航与下单逻辑去堆SEO", "—",
         "无", "不写进合同", "无", "不能承诺"],
    ]
    for r, row in enumerate(rows, 2):
        for c, val in enumerate(row, 1):
            ws.cell(r, c, val)
    style_rows(ws, 1 + len(rows), len(headers), conclusion_col=8)
    for r in range(2, 14):
        ws.row_dimensions[r].height = 78
    set_widths(ws, {"A": 8, "B": 12, "C": 36, "D": 22, "E": 36, "F": 42, "G": 26, "H": 24})
    finish_table(ws, 1 + len(rows), len(headers))
    ws["A15"] = "SEO vs GEO"
    ws["A15"].font = SUB_FONT
    ws["A16"] = (
        "SEO：让Google等搜索引擎看懂页面（title、sitemap、JSON-LD、速度）。"
        "GEO/AEO：让ChatGPT、Perplexity、Google AI Overviews 等生成式引擎在引用时找得到、引用得准（llms.txt、原子事实、AI爬虫放行、知识库与页面一致）。"
        "开源能把「可被理解、可被抓取、不互相矛盾」做完；没有任何开源插件能保证排名或保证被点名引用。"
        "slam.systems 是静态站，Yoast/Rank Math/AEO God Mode 不能直接安装，但同等能力全部可以手写进现有HTML，且不改现有业务逻辑。"
    )
    ws["A16"].alignment = WRAP
    ws["A16"].font = BODY
    ws.row_dimensions[16].height = 90


def plan_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("分期实施计划")
    headers = ["阶段", "目标", "包含的问题/功能", "开源与自研", "依赖", "交付物", "不做"]
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)
    style_header(ws, len(headers), P1)
    rows = [
        ["一期 · P1 拓客闭环", "能按HS/国家产出可导出的清洗后进口商表",
         "Q1–Q6，F1–F4",
         "自研Discovery + 本仓Hunter + crawlee-python + comtradeapicall + 1家商业海关API",
         "客户给出目标国、HS样本、是否采购海关库",
         "任务配置、线索表CSV/JSONL、证据字段、过滤规则开关",
         "不接CRM；不爬付费海关前台"],
        ["二期 · P2 接待 + SEO/GEO", "官网能24h用手册答、不幻觉；页面可被搜索引擎和AI爬虫读懂",
         "Q7–Q12，F5–F9，S1–S10",
         "RAG挂件（LangGraph/LiteLLM）+ 可选Chatwoot；静态站手写JSON-LD/llms.txt/sitemap（aio-surfaces/llms-txt）",
         "产品手册、禁答清单、EU托管选择、是否允许AI爬虫",
         "挂件、知识库后台、WhatsApp深链、SEO/GEO文件、GSC提交记录",
         "不改现有导航/表单；不承诺排名与必被ChatGPT引用"],
        ["三期 · P3 发帖 + 筛客", "OAuth一键分发；关键词+LLM筛客进P1池",
         "Q13–Q22，F10–F11",
         "独立部署Postiz（AGPL隔离）+ 本仓LLM筛客 + 自有频道Captions/Whisper",
         "客户接受不做自动赞粉；准备各平台开发者账号",
         "绑定/解绑、定时发布、筛客列表、审计日志、熔断开关",
         "自动赞/粉、云手机、RPA、封号必解"],
    ]
    for r, row in enumerate(rows, 2):
        for c, val in enumerate(row, 1):
            ws.cell(r, c, val)
        ws.row_dimensions[r].height = 96
    style_rows(ws, 4, len(headers), conclusion_col=None)
    set_widths(ws, {"A": 22, "B": 28, "C": 22, "D": 40, "E": 32, "F": 36, "G": 28})
    finish_table(ws, 4, len(headers))


def gate_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("客户拍板清单")
    headers = ["编号", "必须书面确认的事项", "为什么卡住", "可选答案示例", "不确认的后果"]
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h)
    style_header(ws, len(headers), NO)
    rows = [
        ["G1", "目标采购国清单", "决定买哪家海关库、公开源够不够", "例如：德国、阿联酋、美国", "无法承诺企业级覆盖范围"],
        ["G2", "不买商业海关库时，能否接受没有企业级名单（尤其欧陆）", "欧陆公开数据基本没有进口商名录", "接受 / 批准采购Volza或同等API", "合同不能写「全球公开海关企业名单」"],
        ["G3", "样本HS码 + 要排除的货代词表", "一期规则与试点集", "HS + 产品名 + forwarder/logistics排除词", "清洗误杀/漏杀会高"],
        ["G4", "P2挂静态slam.systems还是另有WordPress", "决定用挂件还是WP插件，SEO插件能不能装", "静态站 / 同时有WP则一并挂", "按静态站做；Yoast类不上"],
        ["G5", "P3是否接受「只发帖+筛客，不自动赞粉」", "自动赞粉无法SLA", "接受降级", "P3互动项从范围删除"],
        ["G6", "数据托管地区（GDPR）", "对话与线索驻留", "EU VPC / 客户自有云", "P2不能上生产对话"],
        ["G7", "AI爬虫：允许被引用，还是拒绝训练", "robots.txt 放行还是Disallow GPTBot等", "允许引用 / 拒绝训练 / 只允许搜索不允许训练", "GEO策略无法定稿"],
    ]
    for r, row in enumerate(rows, 2):
        for c, val in enumerate(row, 1):
            ws.cell(r, c, val)
        ws.row_dimensions[r].height = 56
    style_rows(ws, 1 + len(rows), len(headers), conclusion_col=None)
    set_widths(ws, {"A": 8, "B": 42, "C": 32, "D": 36, "E": 36})
    finish_table(ws, 1 + len(rows), len(headers))


def write_csv_zip(wb: Workbook, zip_path: Path) -> None:
    zip_path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(max_row=ws.max_row, max_col=ws.max_column, values_only=True):
                rows.append(["" if v is None else str(v) for v in row])
            csv_name = f"{ws.title}.csv"
            sio = StringIO()
            writer = csv.writer(sio, lineterminator="\n")
            writer.writerows(rows)
            # utf-8-sig so Microsoft Excel on Windows opens Chinese correctly
            zf.writestr(csv_name, sio.getvalue().encode("utf-8-sig"))
    print(f"wrote {zip_path}")


def write_html(wb: Workbook, html_path: Path) -> None:
    parts = [
        "<!doctype html><html lang=zh><head><meta charset=utf-8>",
        "<title>海外拓客需求问答与解决方案</title>",
        "<style>body{font:14px/1.45 Segoe UI,Calibri,sans-serif;margin:0;background:#f4f1ea}",
        "h1{margin:0;padding:16px 20px;background:#1F3A5F;color:#fff;font-size:20px}",
        "nav{display:flex;flex-wrap:wrap;gap:8px;padding:10px 16px;background:#fff;position:sticky;top:0;border-bottom:1px solid #ddd}",
        "nav a{padding:6px 10px;border-radius:6px;background:#1F3A5F;color:#fff;text-decoration:none;font-size:13px}",
        "section{padding:16px} table{border-collapse:collapse;background:#fff}",
        "th,td{border:1px solid #d9d4c8;padding:8px 10px;vertical-align:top;max-width:420px}",
        "th{background:#1F3A5F;color:#fff} tr:nth-child(even) td{background:#f7f4ef}</style></head><body>",
        "<h1>海外拓客需求问答与解决方案</h1><nav>",
    ]
    for ws in wb.worksheets:
        parts.append(f'<a href="#{escape(ws.title)}">{escape(ws.title)}</a>')
    parts.append("</nav>")
    for ws in wb.worksheets:
        parts.append(f'<section id="{escape(ws.title)}"><h2>{escape(ws.title)}</h2><div style="overflow:auto"><table>')
        for i, row in enumerate(ws.iter_rows(max_row=ws.max_row, max_col=ws.max_column, values_only=True), 1):
            parts.append("<tr>")
            tag = "th" if i == 1 else "td"
            for v in row:
                parts.append(f"<{tag}>{escape(str(v or '')).replace(chr(10), '<br>')}</{tag}>")
            parts.append("</tr>")
        parts.append("</table></div></section>")
    parts.append("</body></html>")
    html_path.parent.mkdir(parents=True, exist_ok=True)
    html_path.write_text("".join(parts), encoding="utf-8")
    print(f"wrote {html_path}")


def copy_to(src: Path, dest: Path) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(src.read_bytes())
    print(f"wrote {dest}")


def main() -> None:
    wb = Workbook()
    wb.properties.creator = "AI Hunter"
    wb.properties.title = "海外拓客需求问答与解决方案"
    cover_sheet(wb)
    qa_sheet(wb)
    features_sheet(wb)
    github_sheet(wb)
    seo_sheet(wb)
    plan_sheet(wb)
    gate_sheet(wb)

    docs = Path("/workspace/docs")
    desktop = Path.home() / "Desktop"
    artifacts = Path("/opt/cursor/artifacts")
    desktop.mkdir(parents=True, exist_ok=True)
    docs.mkdir(parents=True, exist_ok=True)

    cn_name = "海外拓客需求问答与解决方案.xlsx"
    en_name = "overseas-gtm-qa-solutions.xlsx"
    primary = docs / en_name
    wb.save(primary)
    print(f"wrote {primary}")

    for folder in (docs, desktop, artifacts):
        if folder == artifacts and not artifacts.exists():
            continue
        copy_to(primary, folder / cn_name)
        copy_to(primary, folder / en_name)

    write_csv_zip(wb, docs / "overseas-gtm-qa-solutions-csv.zip")
    write_csv_zip(wb, desktop / "overseas-gtm-qa-solutions-csv.zip")
    if artifacts.exists():
        write_csv_zip(wb, artifacts / "overseas-gtm-qa-solutions-csv.zip")

    write_html(wb, docs / "overseas-gtm-qa-solutions.html")
    write_html(wb, desktop / "overseas-gtm-qa-solutions.html")
    if artifacts.exists():
        write_html(wb, artifacts / "overseas-gtm-qa-solutions.html")


if __name__ == "__main__":
    main()
